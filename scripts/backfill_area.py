"""Backfill the `Area` column for rows already present in an existing output spreadsheet.

Reads `Query ID`s with empty `Area` from the configured output file, re-runs the
scraping pipeline (captcha + submit + PDF parse) for each of them, and updates
only the `Area` cell in-place via `SpreadsheetWriter.update_area`.

Usage:
    DOC_OUTPUT_PATH="artifacts/doc_data/araxa_avcb.xlsx" python scripts/backfill_area.py
"""
from __future__ import annotations

import logging
import os
import sys
from dataclasses import replace
from pathlib import Path

import openpyxl

# Force utf-8 on stdout so Playwright's debug logs (containing → ← arrows) don't crash logging on Windows cp1252.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

try:
    import secrets_local  # noqa: F401
except ImportError:
    pass

from config import (
    BrowserConfig,
    CaptchaConfig,
    ConfigError,
    FormConfig,
    RateLimitConfig,
    ScraperConfig,
    StorageConfig,
    validate_configs,
)
from src.doc_scraper.browser import BrowserManager
from src.doc_scraper.captcha import CaptchaSolver
from src.doc_scraper.models import DocRecord
from src.doc_scraper.orchestrator import DocScraper
from src.doc_scraper.storage import SpreadsheetWriter

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s — %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger("backfill_area")


def _ids_missing_area(path: Path, sheet_name: str) -> list[str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return []
        try:
            id_idx = header.index("Query ID")
        except ValueError:
            return []
        try:
            area_idx = header.index("Area")
        except ValueError:
            area_idx = None  # column not yet present — every row needs backfill

        ids: list[str] = []
        for row in rows:
            if not row or row[id_idx] in (None, ""):
                continue
            area_val = row[area_idx] if area_idx is not None and area_idx < len(row) else None
            if area_val in (None, ""):
                ids.append(str(row[id_idx]))
        return ids
    finally:
        workbook.close()


def main() -> int:
    storage_config = replace(StorageConfig(), output_path=os.environ.get("DOC_OUTPUT_PATH", StorageConfig().output_path))
    scraper_config = ScraperConfig()
    captcha_config = CaptchaConfig()
    form_config = FormConfig()

    try:
        validate_configs(scraper_config, captcha_config, form_config)
    except ConfigError as exc:
        logger.error("%s", exc)
        return 2

    path = Path(storage_config.output_path)
    if not path.exists():
        logger.error("Output file not found: %s", path)
        return 2

    ids_to_fill = _ids_missing_area(path, storage_config.sheet_name)
    total = len(ids_to_fill)
    logger.info("Found %d row(s) missing Area in %s", total, path)
    if total == 0:
        return 0

    # Instantiating the writer triggers the header migration that adds the Area column.
    writer = SpreadsheetWriter(storage_config)
    scraper = DocScraper(
        scraper_config=scraper_config,
        browser_config=BrowserConfig(),
        captcha_config=captcha_config,
        storage_config=storage_config,
        rate_limit_config=RateLimitConfig(),
        form_config=form_config,
        _shared_writer=writer,
    )

    updated = not_found = empty_area = failed = 0
    with BrowserManager(BrowserConfig()) as browser:
        solver = CaptchaSolver(captcha_config)
        for i, query_id in enumerate(ids_to_fill, start=1):
            logger.info("[%d/%d] Backfilling %s", i, total, query_id)
            try:
                record = scraper.scrape_one(browser, solver, query_id)
            except Exception:
                failed += 1
                logger.exception("[%s] Crash during scrape, skipping ID", query_id)
                continue
            if not isinstance(record, DocRecord):
                failed += 1
                logger.warning("[%s] No DocRecord returned (got %r)", query_id, type(record).__name__)
                continue
            if not record.area:
                empty_area += 1
                logger.warning("[%s] Scraped successfully but PDF had no Area field", query_id)
                continue
            if writer.update_area(query_id, record.area):
                updated += 1
            else:
                not_found += 1
                logger.warning("[%s] Row not found in spreadsheet, skipping update", query_id)

    logger.info(
        "Backfill complete — updated=%d, empty_area=%d, not_found=%d, failed=%d",
        updated, empty_area, not_found, failed,
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())

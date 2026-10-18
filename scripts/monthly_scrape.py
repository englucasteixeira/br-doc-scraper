"""Monthly incremental scrape: new IDs of the current and previous year, plus a
re-check of known documents that are about to expire.

Steps:
1. Health check: query one known ID; abort with exit 4 if the portal does not answer.
2. For the current and previous year, extend the known range upwards with
   `find_range` (a year with no known range gets a full discovery).
3. Scrape `done_through + 1 .. high` of each year into doc_data_{year}.xlsx
   (resume-safe). `done_through` only advances up to the first failed ID, so a
   failed slice is retried next month.
4. Re-query, by Doc Number, every Records row of --output-dir whose validity is
   between today and today + --revise-days, into doc_data_rev_{YYYY}_{MM}.xlsx.
   Its name sorts after the yearly files, so the pipeline reads it last.

State (--state, JSON): {"2026": {"low": 11, "high": 215031, "done_through": 215031}}.
A missing state file is seeded from --ranges with done_through = high (years
already fully scraped).

Exit codes: 0 complete; 3 partial (some IDs exhausted their attempts, retried
next month); 2 configuration error; 4 portal unreachable or nothing answered.

Usage:
    python scripts/monthly_scrape.py --output-dir ../renovacao-avcb-crm/data/raw \\
        --state ../renovacao-avcb-crm/data/state/scrape_state.json
"""

from __future__ import annotations

import argparse
import json
import logging
import sys
import tempfile
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable, Optional

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

try:
    import secrets_local  # noqa: F401
except ImportError:
    pass

import openpyxl

from config import (
    BrowserConfig,
    CaptchaConfig,
    ConfigError,
    FormConfig,
    RateLimitConfig,
    ScraperConfig,
    StorageConfig,
    validate_configs,
)
from scripts.find_year_ranges import ProbeFn, _make_real_probe, find_range
from src.doc_scraper.browser import BrowserManager
from src.doc_scraper.captcha import CaptchaSolver
from src.doc_scraper.orchestrator import DocScraper, RunStats

logger = logging.getLogger("monthly_scrape")

EXIT_OK = 0
EXIT_CONFIG = 2
EXIT_PARTIAL = 3
EXIT_UNREACHABLE = 4

# Error categories that already settle a document; it is not re-checked again.
SETTLED_CATEGORIES = {"EXPIRED", "REPLACED"}


@dataclass
class YearState:
    low: int
    high: int
    done_through: int

    def to_json(self) -> dict:
        return {"low": self.low, "high": self.high, "done_through": self.done_through}


def load_state(state_path: Path, ranges_path: Path) -> dict[int, YearState]:
    if state_path.exists():
        data = json.loads(state_path.read_text(encoding="utf-8"))
        return {int(year): YearState(**value) for year, value in data.items()}
    if not ranges_path.exists():
        return {}
    data = json.loads(ranges_path.read_text(encoding="utf-8"))
    return {
        int(year): YearState(low=rng[0], high=rng[1], done_through=rng[1])
        for year, rng in data.items()
        if rng is not None
    }


def save_state(state_path: Path, state: dict[int, YearState]) -> None:
    state_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {str(year): state[year].to_json() for year in sorted(state)}
    tmp = state_path.with_suffix(".tmp")
    tmp.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    tmp.replace(state_path)


def years_to_extend(today: date) -> list[int]:
    """New IDs only appear in the current year, and in the previous one around January."""
    return [today.year - 1, today.year]


class AmbiguousProbe(Exception):
    """A probe exhausted its attempts; the range search cannot be trusted."""


def extend_year(
    probe: ProbeFn,
    current: Optional[YearState],
    *,
    max_tail: int = 1_000_000,
    gap_tolerance: int = 5,
) -> Optional[YearState]:
    """Grow a known range upwards, or discover a new year from scratch.

    An ambiguous probe (None) aborts the search and keeps the known range: counting it
    as "exists" would push `high` past the real end, up to max_tail if the portal drops.
    """

    def strict(tail: int) -> Optional[bool]:
        result = probe(tail)
        if result is None:
            raise AmbiguousProbe(tail)
        return result

    try:
        return _extend_year(
            strict, current, max_tail=max_tail, gap_tolerance=gap_tolerance
        )
    except AmbiguousProbe as exc:
        logger.warning("Ambiguous probe at tail %s; range kept as it was.", exc)
        return current


def _extend_year(
    probe: ProbeFn, current: Optional[YearState], *, max_tail: int, gap_tolerance: int
) -> Optional[YearState]:
    if current is None:
        found = find_range(
            probe, max_tail=max_tail, gap_tolerance=gap_tolerance, start_tail=11
        )
        if found is None:
            return None
        low, high = found
        return YearState(low=low, high=high, done_through=low - 1)
    found = find_range(
        probe, max_tail=max_tail, gap_tolerance=gap_tolerance, start_tail=current.high
    )
    high = max(current.high, found[1]) if found is not None else current.high
    return replace(current, high=high)


def advance_done_through(year: YearState, stats: RunStats, prefix: str) -> YearState:
    """`done_through` moves to the last ID known to exist, and never past a failed ID.

    Not-found IDs above the last existing one stay pending, so numbers issued later (or
    a `high` that overshot the real end) are queried again next month."""

    def tails(ids: list[str]) -> list[int]:
        return [int(qid[len(prefix) :]) for qid in ids if qid.startswith(prefix)]

    limit = min(tails(stats.failed), default=year.high + 1) - 1
    existing = [tail for tail in tails(stats.existing) if tail <= limit]
    done = max(existing, default=year.done_through)
    return replace(year, done_through=max(year.done_through, done))


def _parse_validity(value: object) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%d/%m/%Y").date()
        except ValueError:
            return None
    return None


def _header_index(header: tuple, names: tuple[str, ...]) -> Optional[dict[str, int]]:
    cells = [str(cell).strip() if cell is not None else "" for cell in header]
    if not all(name in cells for name in names):
        return None
    return {name: cells.index(name) for name in names}


def select_revisions(
    output_dir: Path, today: date, days: int, limit: Optional[int] = None
) -> list[str]:
    """Doc Numbers in the Records sheets that expire in [today, today + days], skipping
    the ones some Errors sheet already marks as EXPIRED or REPLACED. Sorted by validity."""
    until = today + timedelta(days=days)
    due: dict[str, date] = {}
    settled: set[str] = set()
    for path in sorted(output_dir.glob("doc_data_*.xlsx")):
        workbook = openpyxl.load_workbook(path, read_only=True)
        try:
            if "Records" in workbook.sheetnames:
                rows = workbook["Records"].iter_rows(values_only=True)
                index = _header_index(next(rows, ()), ("Doc Number", "Validity"))
                for row in rows if index else ():
                    doc = row[index["Doc Number"]]
                    validity = _parse_validity(row[index["Validity"]])
                    if not doc:
                        continue
                    # Files are read in name order, so a later revision overrides.
                    if validity is not None and today <= validity <= until:
                        due[str(doc).strip()] = validity
                    else:
                        due.pop(str(doc).strip(), None)
            if "Errors" in workbook.sheetnames:
                rows = workbook["Errors"].iter_rows(values_only=True)
                index = _header_index(next(rows, ()), ("Query ID", "Category"))
                for row in rows if index else ():
                    if (
                        row[index["Category"]] in SETTLED_CATEGORIES
                        and row[index["Query ID"]]
                    ):
                        settled.add(str(row[index["Query ID"]]).strip())
        finally:
            workbook.close()
    selected = sorted((v, doc) for doc, v in due.items() if doc not in settled)
    ids = [doc for _, doc in selected]
    return ids[:limit] if limit is not None else ids


def group_by_prefix(doc_numbers: list[str]) -> dict[str, list[str]]:
    """Group by year prefix (PRJ2024...), keeping the order inside each group."""
    groups: dict[str, list[str]] = {}
    for doc in doc_numbers:
        groups.setdefault(doc[:7], []).append(doc)
    return groups


def exit_code(stats: RunStats, health_ok: bool) -> int:
    if not health_ok:
        return EXIT_UNREACHABLE
    if stats.failed and stats.answered == 0:
        return EXIT_UNREACHABLE
    return EXIT_PARTIAL if stats.failed else EXIT_OK


def health_check(
    probe_factory: Callable[[str], ProbeFn], state: dict[int, YearState]
) -> bool:
    """True when the portal answers (any classified result) for a known ID."""
    if not state:
        return True
    year = max(state)
    return probe_factory(f"PRJ{year}")(state[year].low) is not None


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--state", required=True, type=Path)
    parser.add_argument(
        "--ranges",
        type=Path,
        default=ROOT / "year_ranges.json",
        help="Seeds --state when it does not exist yet.",
    )
    parser.add_argument("--workers", type=int, default=2)
    parser.add_argument("--revise-days", type=int, default=60)
    parser.add_argument("--max-revisions", type=int, default=None)
    parser.add_argument("--max-tail", type=int, default=1_000_000)
    parser.add_argument("--gap-tolerance", type=int, default=5)
    parser.add_argument(
        "--skip-extend",
        action="store_true",
        help="Keep the ranges in --state as they are (rehearsals, or redoing a month).",
    )
    parser.add_argument(
        "--today", type=date.fromisoformat, default=None, help=argparse.SUPPRESS
    )
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(threadName)s — %(message)s",
        handlers=[
            logging.StreamHandler(sys.stdout),
            logging.FileHandler(ROOT / "monthly_scrape.log", encoding="utf-8"),
        ],
    )
    today = args.today or date.today()

    base_scraper = ScraperConfig()
    captcha_config = CaptchaConfig()
    form_config = FormConfig()
    try:
        validate_configs(base_scraper, captcha_config, form_config)
    except ConfigError as exc:
        logger.error("%s", exc)
        return EXIT_CONFIG

    args.output_dir.mkdir(parents=True, exist_ok=True)
    state = load_state(args.state, args.ranges)
    common = dict(
        browser_config=BrowserConfig(),
        captcha_config=captcha_config,
        rate_limit_config=RateLimitConfig(),
        form_config=form_config,
    )
    total = RunStats()

    # 1-2. Health check and range extension share one browser and a throwaway workbook.
    with tempfile.TemporaryDirectory() as tmpdir:
        prober = DocScraper(
            scraper_config=base_scraper,
            storage_config=StorageConfig(output_path=str(Path(tmpdir) / "probe.xlsx")),
            **common,
        )
        with BrowserManager(BrowserConfig()) as browser:
            solver = CaptchaSolver(captcha_config)

            def probe_for(prefix: str) -> ProbeFn:
                return _make_real_probe(
                    prober, browser, solver, prefix, base_scraper.id_zero_padding
                )

            if not health_check(probe_for, state):
                logger.error(
                    "Portal did not answer the health check; nothing was scraped."
                )
                return EXIT_UNREACHABLE
            for year in [] if args.skip_extend else years_to_extend(today):
                extended = extend_year(
                    probe_for(f"PRJ{year}"),
                    state.get(year),
                    max_tail=args.max_tail,
                    gap_tolerance=args.gap_tolerance,
                )
                if extended is None:
                    logger.info("Year %d: no IDs yet.", year)
                    continue
                logger.info(
                    "Year %d: range %d..%d, done through %d.",
                    year,
                    extended.low,
                    extended.high,
                    extended.done_through,
                )
                state[year] = extended
                save_state(args.state, state)

    # 3. New IDs, year by year.
    for year in sorted(state):
        year_state = state[year]
        if year_state.done_through >= year_state.high:
            continue
        prefix = f"PRJ{year}"
        start = year_state.done_through + 1
        logger.info("=== Year %d: IDs %d..%d ===", year, start, year_state.high)
        stats = DocScraper.run_parallel(
            num_workers=args.workers,
            scraper_config=replace(
                base_scraper, id_prefix=prefix, id_start=start, id_end=year_state.high
            ),
            storage_config=StorageConfig(
                output_path=str(args.output_dir / f"doc_data_{year}.xlsx")
            ),
            **common,
        )
        total.merge(stats)
        state[year] = advance_done_through(year_state, stats, prefix)
        save_state(args.state, state)

    # 4. Re-check documents about to expire.
    revisions = select_revisions(
        args.output_dir, today, args.revise_days, args.max_revisions
    )
    logger.info(
        "=== Re-checking %d documents expiring by %s ===",
        len(revisions),
        today + timedelta(days=args.revise_days),
    )
    revision_storage = StorageConfig(
        output_path=str(args.output_dir / f"doc_data_rev_{today:%Y_%m}.xlsx")
    )
    for prefix, ids in group_by_prefix(revisions).items():
        # The year prefix lets a REPLACED document follow its tail to the successor.
        stats = DocScraper.run_parallel(
            num_workers=args.workers,
            scraper_config=replace(base_scraper, id_prefix=prefix),
            storage_config=revision_storage,
            query_ids=ids,
            record_replaced=True,
            **common,
        )
        total.merge(stats)

    code = exit_code(total, health_ok=True)
    logger.info(
        "Summary: %d records, %d errors, %d not found, %d failed; exit %d.",
        total.records,
        total.errors,
        total.not_found,
        len(total.failed),
        code,
    )
    if total.failed:
        logger.warning("Failed IDs (first 20): %s", ", ".join(total.failed[:20]))
    return code


if __name__ == "__main__":
    sys.exit(main())

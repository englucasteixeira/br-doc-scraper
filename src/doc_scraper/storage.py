from __future__ import annotations

import logging
import threading
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.styles import Font

from config import StorageConfig
from .models import DocRecord, ErrorRecord

logger = logging.getLogger(__name__)

_HEADERS = [
    "Query ID", "Doc Number", "Entity Name", "Entity ID",
    "Address", "District", "City", "Validity", "Status",
    "Category", "Audience", "Owner", "Operator",
]


_ERROR_HEADERS = ["Query ID", "Error"]


class SpreadsheetWriter:
    """Persists DocRecord objects to an Excel spreadsheet incrementally."""

    def __init__(self, config: StorageConfig = StorageConfig()) -> None:
        self._config = config
        self._path = Path(config.output_path)
        self._lock = threading.Lock()
        self._workbook, self._sheet, self._error_sheet = self._load_or_create()

    def already_scraped_ids(self) -> set[str]:
        ids: set[str] = set()
        with self._lock:
            for sheet in (self._sheet, self._error_sheet):
                for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True):
                    if row and row[0]:
                        ids.add(str(row[0]))
        return ids

    def append(self, record: DocRecord) -> None:
        with self._lock:
            row = [record.to_row().get(h) for h in _HEADERS]
            self._sheet.append(row)
            self._workbook.save(self._path)
            logger.debug("Record '%s' written to '%s'.", record.query_id, self._path)

    def append_error(self, record: ErrorRecord) -> None:
        with self._lock:
            self._error_sheet.append([record.query_id, record.error_message])
            self._workbook.save(self._path)
            logger.debug("Error record '%s' written to '%s'.", record.query_id, self._path)

    def append_many(self, records: List[DocRecord]) -> None:
        for record in records:
            self.append(record)

    def _load_or_create(self):
        if self._path.exists():
            workbook = openpyxl.load_workbook(self._path)
            sheet = workbook[self._config.sheet_name] if self._config.sheet_name in workbook.sheetnames else workbook.active
            if self._config.error_sheet_name in workbook.sheetnames:
                error_sheet = workbook[self._config.error_sheet_name]
            else:
                error_sheet = workbook.create_sheet(self._config.error_sheet_name)
                self._write_error_header(error_sheet)
                workbook.save(self._path)
        else:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = self._config.sheet_name
            self._write_header(sheet)
            error_sheet = workbook.create_sheet(self._config.error_sheet_name)
            self._write_error_header(error_sheet)
            workbook.save(self._path)
        return workbook, sheet, error_sheet

    @staticmethod
    def _write_header(sheet) -> None:
        sheet.append(_HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)

    @staticmethod
    def _write_error_header(sheet) -> None:
        sheet.append(_ERROR_HEADERS)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
